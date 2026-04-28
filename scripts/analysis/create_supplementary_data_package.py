#!/usr/bin/env python3
"""Create formatted Supplementary Data S2-S5 and Zenodo-ready tarballs."""

from __future__ import annotations

import argparse
import gzip
import hashlib
import os
import re
import shutil
import subprocess
import tarfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

import pandas as pd


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_MTIME = int(datetime(2026, 4, 18, tzinfo=timezone.utc).timestamp())


@dataclass(frozen=True)
class SheetSpec:
    sheet_name: str
    tsv_stem: str
    frame: pd.DataFrame
    source: str
    description: str


@dataclass(frozen=True)
class DataObject:
    object_id: str
    workbook_name: str
    description: str
    sheets: list[SheetSpec]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Create formatted Supplementary Data S2-S5 and Zenodo tarballs."
    )
    parser.add_argument("--intermediate-data", default="intermediate_data_whole_genome")
    parser.add_argument("--figures-dir", default="figures_whole_genome")
    parser.add_argument("--tables-dir", default="tables_whole_genome")
    parser.add_argument("--output-tables-dir", default=None)
    parser.add_argument("--staging-dir", default=".zenodo_staging")
    parser.add_argument("--dist-dir", default="zenodo_dist")
    parser.add_argument("--version", default="v1.0")
    parser.add_argument(
        "--skip-tarballs",
        action="store_true",
        help="Create data files and staging directories, but do not create tarballs.",
    )
    return parser.parse_args()


def rel(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(REPO_ROOT))
    except ValueError:
        return str(path)


def require_file(path: Path) -> Path:
    if not path.exists():
        raise FileNotFoundError(f"Required input is missing: {rel(path)}")
    return path


def read_tsv(path: Path) -> pd.DataFrame:
    return pd.read_csv(require_file(path), sep="\t")


def read_csv(path: Path) -> pd.DataFrame:
    return pd.read_csv(require_file(path))


def read_parquet(path: Path) -> pd.DataFrame:
    return pd.read_parquet(require_file(path))


def sort_frame(df: pd.DataFrame, preferred: Iterable[str]) -> pd.DataFrame:
    cols = [col for col in preferred if col in df.columns]
    if not cols:
        return df.reset_index(drop=True)
    return df.sort_values(cols, kind="mergesort").reset_index(drop=True)


def normalize_unique_status(series: pd.Series) -> pd.Series:
    return series.astype(str).replace(
        {
            "Syntenic": "Unique",
            "Nonsyntenic": "Nonunique",
            "All": "All",
            "nan": pd.NA,
        }
    )


def derive_file_unique_status(files: pd.Series) -> pd.Series:
    text = files.fillna("").astype(str).str.lower()
    status = pd.Series("All", index=files.index, dtype="object")
    status.loc[text.str.contains("nonsyntenic", regex=False)] = "Nonunique"
    status.loc[
        text.str.contains("syntenic", regex=False)
        & ~text.str.contains("nonsyntenic", regex=False)
    ] = "Unique"
    return status


def clean_panel_label(panel: pd.Series) -> pd.Series:
    panel_base = (
        panel.astype(str)
        .str.replace(r"\.GWAS_filtered\.info_cutoff_0\.\d+$", "", regex=True)
        .str.replace(".common_variants", "", regex=False)
    )
    return panel_base.replace({"native_panel": "Native", "lifted_panel": "Lifted"})


def add_imputation_labels(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out = out.replace({"T2T": "CHM13v2.0"})

    panel_raw = out["panel"].astype(str)
    panel_clean = clean_panel_label(panel_raw)
    out["panel_type"] = panel_clean
    out["variant_subset"] = "All panel variants"
    out.loc[panel_raw.str.contains("common_variants", regex=False), "variant_subset"] = (
        "Common variants"
    )
    out["info_cutoff"] = panel_raw.str.extract(r"info_cutoff_(\d+\.\d+)", expand=False)
    out["info_cutoff"] = out["info_cutoff"].fillna("0")

    out["reference_panel"] = "GRCh38"
    out.loc[
        ((out["genome"] == "GRCh38") & (panel_clean == "Lifted"))
        | ((out["genome"] == "CHM13v2.0") & (panel_clean == "Native")),
        "reference_panel",
    ] = "CHM13v2.0"

    dataset_raw = out["dataset"].astype(str)
    out["panel_filter"] = (
        dataset_raw.str.replace(r"^(SGDP|pangenome)", "", regex=True)
        .str.replace(r"^_", "", regex=True)
        .replace("", "No Filter")
    )
    out["dataset"] = dataset_raw.str.split("_", n=1).str[0].replace(
        {"pangenome": "Pangenome"}
    )

    if "Synteny" in out.columns:
        out["synteny_status"] = out["Synteny"]
        out["unique_status"] = normalize_unique_status(out["Synteny"])
    elif "file" in out.columns:
        out["unique_status"] = derive_file_unique_status(out["file"])
        out["synteny_status"] = out["unique_status"].replace(
            {"Unique": "Syntenic", "Nonunique": "Nonsyntenic"}
        )
    else:
        out["unique_status"] = "All"
        out["synteny_status"] = "All"

    if "var_category_id" in out.columns:
        out = out.rename(columns={"var_category_id": "variant_category"})

    if {"num_Aa_mismatches", "num_aa_mismatches"}.issubset(out.columns):
        out["num_alt_mismatches"] = out["num_Aa_mismatches"] + out["num_aa_mismatches"]
    if {"num_Aa", "num_aa"}.issubset(out.columns):
        out["num_alt_variants"] = out["num_Aa"] + out["num_aa"]
    if {"num_alt_mismatches", "num_alt_variants"}.issubset(out.columns):
        out["non_reference_discordance_percent_recomputed"] = (
            out["num_alt_mismatches"] / out["num_alt_variants"] * 100
        )

    if "file" in out.columns:
        out = out.rename(columns={"file": "source_file"})
    if "genome" in out.columns:
        out = out.rename(columns={"genome": "target_genome"})

    leading = [
        "dataset",
        "ancestry",
        "sample_name",
        "variant_category",
        "target_genome",
        "reference_panel",
        "panel_type",
        "variant_subset",
        "panel_filter",
        "info_cutoff",
        "unique_status",
        "synteny_status",
        "variant_bin",
        "mean_AF",
    ]
    ordered = [col for col in leading if col in out.columns] + [
        col for col in out.columns if col not in leading
    ]
    return out[ordered]


def build_s2(output_tables_dir: Path) -> DataObject:
    per_sample = read_tsv(output_tables_dir / "supplemental_per_sample_data_unformatted.tsv")
    source_summary = read_tsv(
        output_tables_dir / "supplemental_ground_truth_source_accuracy_data_unformatted.tsv"
    )
    per_sample = sort_frame(
        per_sample,
        [
            "panel_genome",
            "ground_truth_data_source",
            "family_status",
            "superpopulation",
            "population",
            "sample_id",
        ],
    )
    source_summary = sort_frame(
        source_summary, ["panel_genome", "ground_truth_data_source", "family_status"]
    )
    return DataObject(
        object_id="S2",
        workbook_name="Supplementary_Data_S2.xlsx",
        description=(
            "Per-sample phasing and genotyping accuracy resource, stratified by "
            "reference panel, phasing method, and ground-truth data source."
        ),
        sheets=[
            SheetSpec(
                "per_sample_accuracy",
                "Supplementary_Data_S2_per_sample_accuracy",
                per_sample,
                "supplemental_per_sample_data_unformatted.tsv",
                "Per-sample variant counts, error counts, switch error rates, and genotyping error rates.",
            ),
            SheetSpec(
                "ground_truth_source_summary",
                "Supplementary_Data_S2_ground_truth_source_summary",
                source_summary,
                "supplemental_ground_truth_source_accuracy_data_unformatted.tsv",
                "Accuracy summary stratified by panel, phasing method, and ground-truth data source.",
            ),
        ],
    )


def build_s3(intermediate_dir: Path) -> DataObject:
    source_name = "per_MAF_bin.parquet"
    source_path = intermediate_dir / source_name
    if not source_path.exists():
        source_name = "binned_maf_data.parquet"
        source_path = intermediate_dir / source_name
    maf = read_parquet(source_path)
    maf = maf.copy()
    if "type" in maf.columns:
        maf = maf.rename(columns={"type": "variant_type"})
    if "genome" in maf.columns:
        maf = maf.rename(columns={"genome": "panel_genome"})
    if "syntenic" in maf.columns:
        maf["synteny_status"] = maf["syntenic"]
        maf["unique_status"] = normalize_unique_status(maf["syntenic"])
        maf = maf.drop(columns=["syntenic"])
    elif "region" in maf.columns:
        maf["synteny_status"] = maf["region"].where(
            maf["region"].isin(["Syntenic", "Nonsyntenic", "All"]), "All"
        )
        maf["unique_status"] = normalize_unique_status(maf["synteny_status"])
    maf = sort_frame(
        maf,
        [
            "panel_genome",
            "method_of_phasing",
            "ground_truth_data_source",
            "unique_status",
            "variant_type",
            "rounded_MAF",
        ],
    )
    return DataObject(
        object_id="S3",
        workbook_name="Supplementary_Data_S3.xlsx",
        description=(
            "Minor-allele-frequency-bin phasing and genotyping accuracy resource, "
            "stratified by panel, phasing method, unique/nonunique status, and "
            "ground-truth data source."
        ),
        sheets=[
            SheetSpec(
                "maf_bin_accuracy",
                "Supplementary_Data_S3_maf_bin_accuracy",
                maf,
                source_name,
                "Per-MAF-bin variant counts, error counts, switch error rates, and genotyping error rates.",
            )
        ],
    )


def build_s4(intermediate_dir: Path) -> DataObject:
    per_sample = add_imputation_labels(
        read_parquet(intermediate_dir / "per_sample_imputation_performance.parquet")
    )
    per_variant = add_imputation_labels(
        read_parquet(intermediate_dir / "per_variant_category_imputation_performance.parquet")
    )
    per_sample = sort_frame(
        per_sample,
        [
            "dataset",
            "ancestry",
            "sample_name",
            "target_genome",
            "reference_panel",
            "panel_type",
            "variant_subset",
            "panel_filter",
            "unique_status",
            "variant_category",
        ],
    )
    per_variant = sort_frame(
        per_variant,
        [
            "dataset",
            "ancestry",
            "target_genome",
            "reference_panel",
            "panel_type",
            "variant_subset",
            "panel_filter",
            "unique_status",
            "variant_category",
            "variant_bin",
            "mean_AF",
        ],
    )
    return DataObject(
        object_id="S4",
        workbook_name="Supplementary_Data_S4.xlsx",
        description=(
            "Imputation performance resource for sample-level and variant-category "
            "benchmarks across panels, datasets, ancestry groups, and "
            "unique/nonunique status."
        ),
        sheets=[
            SheetSpec(
                "per_sample_imputation",
                "Supplementary_Data_S4_per_sample_imputation",
                per_sample,
                "per_sample_imputation_performance.parquet",
                "Per-subject imputation summary statistics by panel, variant category, and unique/nonunique status.",
            ),
            SheetSpec(
                "variant_category_imputation",
                "Supplementary_Data_S4_variant_category_imputation",
                per_variant,
                "per_variant_category_imputation_performance.parquet",
                "Imputation summary statistics by variant category, panel, ancestry, and unique/nonunique status.",
            ),
        ],
    )


def load_optional_tsv(path: Path) -> pd.DataFrame | None:
    if path.exists():
        return pd.read_csv(path, sep="\t")
    return None


def build_s5(tables_dir: Path) -> DataObject:
    sheet_inputs = [
        (
            "cytoband_ranked_performance",
            "Supplementary_Data_S5_cytoband_ranked_performance",
            "1kgp_variation_HPRC_ground_truth.tsv",
            "Cytobands ranked by T2T-CHM13 versus GRCh38 phasing/genotyping performance.",
            ["rel_drop_SER", "abs_diff_SER", "interval"],
            False,
        ),
        (
            "cytoband_summary",
            "Supplementary_Data_S5_cytoband_summary",
            "1kgp_variation_HPRC_ground_truth_summarized.tsv",
            "Cytoband-level summary of T2T-CHM13 versus GRCh38 performance.",
            ["rel_drop_SER", "abs_diff_SER", "interval"],
            False,
        ),
        (
            "decipher_cnv_regions",
            "Supplementary_Data_S5_decipher_cnv_regions",
            "decipher_cnv_region_performances_rephased_1KGP_variants_HPRC_ground_truth.tsv",
            "Performance within DECIPHER disease-CNV regions.",
            ["rel_drop_SER", "Syndrome"],
            False,
        ),
        (
            "decipher_cnv_regions_pm1mb",
            "Supplementary_Data_S5_decipher_cnv_regions_pm1mb",
            "decipher_cnv_region_performances_rephased_1KGP_variants_HPRC_ground_truth_plusminus_1mb.tsv",
            "Performance within DECIPHER disease-CNV regions plus and minus 1 Mb.",
            ["rel_drop_SER", "Syndrome"],
            False,
        ),
        (
            "decipher_top_regions",
            "Supplementary_Data_S5_decipher_top_regions",
            "decipher_top10_cytoband_summary.tsv",
            "Summary of top-performing cytobands and their overlap with DECIPHER regions.",
            ["rank_metric"],
            True,
        ),
        (
            "decipher_improvement_conc",
            "Supplementary_Data_S5_decipher_improvement_concentration",
            "decipher_cytoband_improvement_concentration.tsv",
            "Permutation analysis of DECIPHER-region concentration among improved cytobands.",
            ["analysis"],
            True,
        ),
        (
            "decipher_delta_ser_perm",
            "Supplementary_Data_S5_decipher_delta_ser_permutation",
            "decipher_cytoband_delta_ser_permutation.tsv",
            "Permutation analysis comparing DECIPHER and non-DECIPHER cytoband SER differences.",
            ["analysis"],
            True,
        ),
        (
            "decipher_enrichment",
            "Supplementary_Data_S5_decipher_enrichment",
            "decipher_cytoband_enrichment.tsv",
            "DECIPHER cytoband enrichment summary, when regenerated by the plotting notebook.",
            ["analysis"],
            True,
        ),
    ]

    sheets: list[SheetSpec] = []
    for sheet_name, stem, filename, description, sort_cols, optional in sheet_inputs:
        path = tables_dir / filename
        frame = load_optional_tsv(path)
        if frame is None:
            if optional:
                continue
            raise FileNotFoundError(f"Required S5 input is missing: {rel(path)}")
        ascending = [False if col in {"rel_drop_SER", "abs_diff_SER"} else True for col in sort_cols if col in frame.columns]
        sort_by = [col for col in sort_cols if col in frame.columns]
        if sort_by:
            frame = frame.sort_values(sort_by, ascending=ascending, kind="mergesort")
        frame = frame.reset_index(drop=True)
        sheets.append(SheetSpec(sheet_name, stem, frame, filename, description))

    return DataObject(
        object_id="S5",
        workbook_name="Supplementary_Data_S5.xlsx",
        description=(
            "Cytoband and disease-CNV-region performance resource for comparing "
            "GRCh38 and T2T-CHM13 panel behavior in reusable genomic regions."
        ),
        sheets=sheets,
    )


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def pick_excel_engine() -> str:
    try:
        import xlsxwriter  # noqa: F401

        return "xlsxwriter"
    except ImportError:
        return "openpyxl"


def width_for_series(series: pd.Series, column: str) -> int:
    sample = series.dropna().astype(str).head(200)
    max_value = max([len(column), *(sample.map(len).tolist() or [0])])
    return min(max(max_value + 2, 10), 48)


def format_xlsxwriter(writer: pd.ExcelWriter, sheet: str, df: pd.DataFrame) -> None:
    workbook = writer.book
    worksheet = writer.sheets[sheet]
    header_fmt = workbook.add_format(
        {"bold": True, "text_wrap": True, "valign": "top", "border": 1}
    )
    for idx, col in enumerate(df.columns):
        worksheet.write(0, idx, col, header_fmt)
        worksheet.set_column(idx, idx, width_for_series(df[col], col))
    if len(df.columns):
        worksheet.autofilter(0, 0, max(len(df), 1), len(df.columns) - 1)
    worksheet.freeze_panes(1, 0)


def format_openpyxl(writer: pd.ExcelWriter, sheet: str, df: pd.DataFrame) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    worksheet = writer.sheets[sheet]
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for idx, col in enumerate(df.columns, start=1):
        letter = worksheet.cell(row=1, column=idx).column_letter
        worksheet.column_dimensions[letter].width = width_for_series(df[col], col)


def build_workbook_dictionary(obj: DataObject) -> pd.DataFrame:
    rows: list[dict[str, object]] = [
        {
            "section": "workbook",
            "sheet_name": "",
            "column_name": "",
            "tsv_file": "",
            "rows": "",
            "columns": "",
            "source": "",
            "description": obj.description,
        }
    ]
    rows.extend(
        {
            "section": "sheet",
            "sheet_name": spec.sheet_name,
            "column_name": "",
            "tsv_file": f"{spec.tsv_stem}.tsv.gz",
            "rows": len(spec.frame),
            "columns": len(spec.frame.columns),
            "source": spec.source,
            "description": spec.description,
        }
        for spec in obj.sheets
    )
    rows.extend(
        {
            "section": "column",
            "sheet_name": spec.sheet_name,
            "column_name": column,
            "tsv_file": "",
            "rows": "",
            "columns": "",
            "source": spec.source,
            "description": COLUMN_DESCRIPTIONS.get(
                column,
                "Source-derived column; see DATA_SOURCES.md and source table provenance.",
            ),
        }
        for spec in obj.sheets
        for column in spec.frame.columns
    )
    return pd.DataFrame(rows)


def write_workbook(path: Path, obj: DataObject) -> None:
    engine = pick_excel_engine()
    with pd.ExcelWriter(path, engine=engine) as writer:
        dictionary_frame = build_workbook_dictionary(obj)
        dictionary_frame.to_excel(writer, sheet_name="dictionary", index=False)
        if engine == "xlsxwriter":
            format_xlsxwriter(writer, "dictionary", dictionary_frame)
        else:
            format_openpyxl(writer, "dictionary", dictionary_frame)
        for spec in obj.sheets:
            spec.frame.to_excel(writer, sheet_name=spec.sheet_name, index=False)
            if engine == "xlsxwriter":
                format_xlsxwriter(writer, spec.sheet_name, spec.frame)
            else:
                format_openpyxl(writer, spec.sheet_name, spec.frame)


def write_tsv_gz(path: Path, df: pd.DataFrame) -> None:
    df.to_csv(
        path,
        sep="\t",
        index=False,
        compression={"method": "gzip", "mtime": 0},
    )


def write_data_objects(objects: list[DataObject], output_tables_dir: Path) -> pd.DataFrame:
    output_tables_dir.mkdir(parents=True, exist_ok=True)
    rows: list[dict[str, object]] = []
    for obj in objects:
        workbook_path = output_tables_dir / obj.workbook_name
        write_workbook(workbook_path, obj)
        rows.append(
            {
                "object": obj.object_id,
                "sheet": "workbook",
                "file": workbook_path.name,
                "format": "xlsx",
                "rows": "",
                "columns": "",
                "sha256": sha256_file(workbook_path),
                "source": "; ".join(spec.source for spec in obj.sheets),
                "description": f"Formatted workbook for Supplementary Data {obj.object_id}.",
            }
        )
        for spec in obj.sheets:
            tsv_path = output_tables_dir / f"{spec.tsv_stem}.tsv.gz"
            write_tsv_gz(tsv_path, spec.frame)
            rows.append(
                {
                    "object": obj.object_id,
                    "sheet": spec.sheet_name,
                    "file": tsv_path.name,
                    "format": "tsv.gz",
                    "rows": len(spec.frame),
                    "columns": len(spec.frame.columns),
                    "sha256": sha256_file(tsv_path),
                    "source": spec.source,
                    "description": spec.description,
                }
            )
    manifest = pd.DataFrame(rows)
    manifest_path = output_tables_dir / "Supplementary_Data_manifest.tsv"
    manifest.to_csv(manifest_path, sep="\t", index=False)
    return manifest


COLUMN_DESCRIPTIONS = {
    "sample_id": "1000 Genomes Project sample identifier.",
    "sample_name": "Sample identifier in the imputation benchmark.",
    "panel_genome": "Genome assembly used for the haplotype panel.",
    "target_genome": "Genome assembly used for the imputation target/evaluation.",
    "reference_panel": "Reference haplotype panel represented by the benchmark row.",
    "panel_type": "Native or lifted panel representation.",
    "variant_subset": "Whether all panel variants or a restricted subset was used.",
    "panel_filter": "Additional panel/data filter, such as no_singletons, snps, or No Filter.",
    "info_cutoff": "Imputation INFO-score cutoff encoded in the source panel name.",
    "ground_truth_data_source": "Assembly or sample set used as ground truth.",
    "family_status": "Trio/proband/parent status used for phasing benchmarks.",
    "population": "1000 Genomes population label.",
    "superpopulation": "1000 Genomes superpopulation label.",
    "sex": "Sample sex annotation.",
    "variant_type": "Variant class used for MAF-bin phasing/genotyping benchmarks.",
    "variant_category": "Variant class used for imputation benchmarks.",
    "unique_status": "Unique/Nonunique/All status derived from synteny categories.",
    "synteny_status": "Syntenic/Nonsyntenic/All status in the source data.",
    "rounded_MAF": "Rounded minor allele frequency bin.",
    "MAF": "Minor allele frequency.",
    "mean_AF": "Mean allele frequency for an imputation variant bin.",
    "variant_bin": "Variant-frequency or imputation bin label from GLIMPSE2 concordance output.",
    "n_switch_errors": "Number of switch errors.",
    "n_checked": "Number of informative phased heterozygous sites checked.",
    "switch_error_rate": "Switch error rate in percent, as represented by the source table.",
    "n_gt_errors": "Number of genotype discordances.",
    "n_gt_checked": "Number of genotypes checked.",
    "gt_error_rate": "Genotype error rate in percent, as represented by the source table.",
    "N50": "Switch-error block N50 from the source phasing benchmark.",
    "best_gt_rsquared": "Best-guess genotype r-squared from imputation evaluation.",
    "imputed_ds_rsquared": "Dosage r-squared from imputation evaluation.",
    "non_reference_discordance_percent": "Non-reference genotype discordance percentage from source data.",
    "rel_drop_SER": "Relative reduction in switch error rate comparing GRCh38 and CHM13 rows.",
    "abs_diff_SER": "Absolute switch error-rate difference comparing GRCh38 and CHM13 rows.",
    "rel_drop_GER": "Relative reduction in genotype error rate comparing GRCh38 and CHM13 rows.",
    "abs_diff_GER": "Absolute genotype error-rate difference comparing GRCh38 and CHM13 rows.",
    "Syndrome": "DECIPHER CNV syndrome or disease-region label.",
    "interval": "Cytoband or genomic interval label.",
}


def write_column_dictionary(objects: list[DataObject], output_tables_dir: Path) -> Path:
    rows = []
    for obj in objects:
        for spec in obj.sheets:
            for column in spec.frame.columns:
                rows.append(
                    {
                        "object": obj.object_id,
                        "sheet": spec.sheet_name,
                        "column": column,
                        "description": COLUMN_DESCRIPTIONS.get(
                            column,
                            "Source-derived column; see DATA_SOURCES.md and source table provenance.",
                        ),
                    }
                )
    dictionary = pd.DataFrame(rows).drop_duplicates()
    path = output_tables_dir / "Supplementary_Data_column_dictionary.tsv"
    dictionary.to_csv(path, sep="\t", index=False)
    return path


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def link_or_copy(src: Path, dst: Path) -> None:
    dst.parent.mkdir(parents=True, exist_ok=True)
    if dst.exists() or dst.is_symlink():
        dst.unlink()
    try:
        dst.symlink_to(src.resolve())
    except OSError:
        shutil.copy2(src, dst)


def package_manifest(root: Path, checksum_name: str = "CHECKSUMS.sha256") -> pd.DataFrame:
    rows = []
    for path in sorted(root.rglob("*")):
        if not path.is_file() and not path.is_symlink():
            continue
        rel_path = path.relative_to(root)
        if rel_path.name == checksum_name:
            continue
        target = path.resolve() if path.is_symlink() else path
        rows.append(
            {
                "path": str(rel_path),
                "bytes": target.stat().st_size,
                "sha256": sha256_file(target),
            }
        )
    return pd.DataFrame(rows)


def write_package_metadata(root: Path, readme: str, data_sources: str) -> None:
    write_text(root / "README.md", readme)
    write_text(root / "DATA_SOURCES.md", data_sources)
    manifest_path = root / "MANIFEST.tsv"
    package_manifest(root).to_csv(manifest_path, sep="\t", index=False)
    manifest = package_manifest(root)
    checksum_lines = [
        f"{row.sha256}  {row.path}" for row in manifest.itertuples(index=False)
    ]
    write_text(root / "CHECKSUMS.sha256", "\n".join(checksum_lines) + "\n")


def provenance_text(args: argparse.Namespace) -> str:
    return (
        "python scripts/analysis/create_supplementary_data_package.py "
        f"--intermediate-data {args.intermediate_data} "
        f"--figures-dir {args.figures_dir} "
        f"--tables-dir {args.tables_dir} "
        f"--version {args.version}\n"
    )


def git_commit_text() -> str:
    try:
        commit = subprocess.check_output(
            ["git", "rev-parse", "HEAD"], cwd=REPO_ROOT, text=True
        ).strip()
        status = subprocess.check_output(
            ["git", "status", "--short"], cwd=REPO_ROOT, text=True
        ).strip()
        return f"commit\t{commit}\nworking_tree_status\n{status}\n"
    except Exception as exc:  # pragma: no cover - provenance best effort
        return f"Unable to capture git metadata: {exc}\n"


def python_packages_text() -> str:
    try:
        return subprocess.check_output(
            ["python", "-m", "pip", "freeze"], cwd=REPO_ROOT, text=True
        )
    except Exception as exc:  # pragma: no cover - provenance best effort
        return f"Unable to capture Python package metadata: {exc}\n"


def data_sources_text() -> str:
    return """# Data Sources

Primary data and code links associated with the preprint:

- bioRxiv preprint: https://www.biorxiv.org/content/10.1101/2025.02.24.639687v1
- Code and performance data Zenodo record: https://zenodo.org/records/12121420
- T2T-CHM13 recombination maps Zenodo record: https://zenodo.org/records/14891074
- GitHub repository: https://github.com/JosephLalli/phasing_T2T
- Public HPRC panel path from the preprint Data/Code tab:
  https://s3-us-west-2.amazonaws.com/human-pangenomics/index.html?prefix=T2T/CHM13/assemblies/variants/1000_Genomes_Project/chm13v2.0/Phased_SHAPEIT5_v1.1/

Raw VCFs, assemblies, large intermediate parquet bundles, and recombination-map archives are not duplicated here when they already have stable public locations.
"""


def create_supplement_package(
    staging_dir: Path,
    output_tables_dir: Path,
    generated_files: Iterable[str],
    args: argparse.Namespace,
) -> Path:
    root = staging_dir / f"phasing_T2T_supplementary_data_S2-S5_{args.version}"
    if root.exists():
        shutil.rmtree(root)
    for filename in generated_files:
        if filename.startswith("Supplementary_Data_S"):
            object_id = filename.split("_")[2].split(".")[0]
            link_or_copy(output_tables_dir / filename, root / "supplementary_data" / object_id / filename)
        elif filename.startswith("Supplementary_Data_"):
            link_or_copy(output_tables_dir / filename, root / filename)

    write_text(root / "provenance" / "generation_command.txt", provenance_text(args))
    write_text(root / "provenance" / "git_commit.txt", git_commit_text())
    write_text(root / "provenance" / "python_packages.txt", python_packages_text())
    write_package_metadata(
        root,
        readme=(
            "# phasing_T2T Supplementary Data S2-S5\n\n"
            "This archive contains formatted XLSX workbooks and machine-readable TSV.GZ "
            "siblings for Supplementary Data S2-S5. Supplementary Data S1 is generated "
            "by the recombination-map workflow and is referenced in DATA_SOURCES.md.\n"
        ),
        data_sources=data_sources_text(),
    )
    return root


def create_source_package(
    staging_dir: Path,
    figures_dir: Path,
    tables_dir: Path,
    args: argparse.Namespace,
) -> Path:
    root = staging_dir / f"phasing_T2T_paper_source_data_{args.version}"
    if root.exists():
        shutil.rmtree(root)

    source_files: list[Path] = []
    source_files.extend(sorted(tables_dir.glob("*.tsv")))
    source_files.extend(sorted(figures_dir.glob("*.csv")))
    source_files.extend(sorted((figures_dir / "supplemental").glob("*.csv")))
    source_files.extend(sorted((figures_dir / "supplemental" / "tables").glob("*.tsv")))
    source_files.extend(sorted((figures_dir / "supplemental" / "tables").glob("Supplementary_Data_*")))

    for src in source_files:
        try:
            relative = src.relative_to(REPO_ROOT)
        except ValueError:
            relative = Path(src.name)
        link_or_copy(src, root / "source_data" / relative)

    write_text(root / "provenance" / "generation_command.txt", provenance_text(args))
    write_text(root / "provenance" / "git_commit.txt", git_commit_text())
    write_text(root / "provenance" / "python_packages.txt", python_packages_text())
    write_package_metadata(
        root,
        readme=(
            "# phasing_T2T Paper Source Data\n\n"
            "This archive contains compact source tables, figure-source CSV/TSV files, "
            "formatted supplementary data, and provenance files for the phasing_T2T paper. "
            "Large upstream resources are referenced in DATA_SOURCES.md rather than duplicated.\n"
        ),
        data_sources=data_sources_text(),
    )
    return root


def normalize_tarinfo(tarinfo: tarfile.TarInfo) -> tarfile.TarInfo:
    tarinfo.uid = 0
    tarinfo.gid = 0
    tarinfo.uname = ""
    tarinfo.gname = ""
    tarinfo.mtime = DEFAULT_MTIME
    return tarinfo


def create_tarball(root: Path, dist_dir: Path) -> Path:
    dist_dir.mkdir(parents=True, exist_ok=True)
    tarball = dist_dir / f"{root.name}.tar.gz"
    with tarball.open("wb") as raw:
        with gzip.GzipFile(fileobj=raw, mode="wb", mtime=0) as gz:
            with tarfile.open(fileobj=gz, mode="w") as tar:
                for path in sorted(root.rglob("*"), key=lambda p: str(p.relative_to(root))):
                    if path.is_dir():
                        continue
                    source = path.resolve() if path.is_symlink() else path
                    arcname = Path(root.name) / path.relative_to(root)
                    tar.add(source, arcname=str(arcname), recursive=False, filter=normalize_tarinfo)
    return tarball


def main() -> int:
    args = parse_args()
    intermediate_dir = (REPO_ROOT / args.intermediate_data).resolve()
    figures_dir = (REPO_ROOT / args.figures_dir).resolve()
    tables_dir = (REPO_ROOT / args.tables_dir).resolve()
    output_tables_dir = (
        (REPO_ROOT / args.output_tables_dir).resolve()
        if args.output_tables_dir
        else figures_dir / "supplemental" / "tables"
    )
    staging_dir = (REPO_ROOT / args.staging_dir).resolve()
    dist_dir = (REPO_ROOT / args.dist_dir).resolve()

    objects = [
        build_s2(output_tables_dir),
        build_s3(intermediate_dir),
        build_s4(intermediate_dir),
        build_s5(tables_dir),
    ]
    manifest = write_data_objects(objects, output_tables_dir)
    dictionary_path = write_column_dictionary(objects, output_tables_dir)
    manifest_path = output_tables_dir / "Supplementary_Data_manifest.tsv"
    manifest.loc[len(manifest)] = {
        "object": "all",
        "sheet": "column_dictionary",
        "file": dictionary_path.name,
        "format": "tsv",
        "rows": len(pd.read_csv(dictionary_path, sep="\t")),
        "columns": 4,
        "sha256": sha256_file(dictionary_path),
        "source": "create_supplementary_data_package.py",
        "description": "Column dictionary for generated Supplementary Data files.",
    }
    manifest.to_csv(manifest_path, sep="\t", index=False)

    generated_files = set(manifest["file"].astype(str))
    generated_files.add(manifest_path.name)
    supplement_root = create_supplement_package(
        staging_dir, output_tables_dir, sorted(generated_files), args
    )
    source_root = create_source_package(staging_dir, figures_dir, tables_dir, args)

    tarballs: list[Path] = []
    if not args.skip_tarballs:
        tarballs = [
            create_tarball(supplement_root, dist_dir),
            create_tarball(source_root, dist_dir),
        ]

    print(f"Wrote supplementary tables to {rel(output_tables_dir)}")
    print(f"Wrote staging directories to {rel(staging_dir)}")
    for tarball in tarballs:
        print(f"Wrote tarball {rel(tarball)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
